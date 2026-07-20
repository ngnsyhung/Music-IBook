from pathlib import Path
from datetime import date
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "Tai_lieu_quan_ly_du_an_Music_IBook.docx"
ASSETS = ROOT / "_project_doc_assets"
ASSETS.mkdir(exist_ok=True)

NAVY = "163A5F"
BLUE = "2E74B5"
LIGHT_BLUE = "E8EEF5"
PALE_BLUE = "F4F8FC"
GOLD = "D9A441"
GRAY = "5B6573"
LIGHT_GRAY = "F2F4F7"
WHITE = "FFFFFF"
RED = "B42318"
GREEN = "18794E"


def set_font(run, size=11, bold=False, color="222222", italic=False, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths):
    total = sum(widths)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths[min(idx, len(widths) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_text(cell, value, bold=False, color="222222", size=8.7, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    r = p.add_run(str(value))
    set_font(r, size=size, bold=bold, color=color)


def add_table(doc, headers, rows, widths, font_size=8.7):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        shade(hdr.cells[i], NAVY)
        set_cell_text(hdr.cells[i], h, bold=True, color=WHITE, size=font_size, align=WD_ALIGN_PARAGRAPH.CENTER)
    for ridx, values in enumerate(rows):
        cells = table.add_row().cells
        for i, v in enumerate(values):
            if ridx % 2 == 1:
                shade(cells[i], PALE_BLUE)
            align = WD_ALIGN_PARAGRAPH.CENTER if i == 0 and len(widths) > 2 else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(cells[i], v, size=font_size, align=align)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    set_font(r, size=10.5)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_font(r, size=10.5)
    return p


def add_callout(doc, title, text, fill=LIGHT_BLUE, accent=BLUE):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    shade(cell, fill)
    set_cell_margins(cell, top=150, bottom=150, start=190, end=190)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title.upper())
    set_font(r, size=9.5, bold=True, color=accent)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.15
    r2 = p2.add_run(text)
    set_font(r2, size=10.2)
    set_table_geometry(table, [9360])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    return p


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.18
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_font(r1, size=10.7, bold=True, color=NAVY)
        r2 = p.add_run(text[len(bold_prefix):])
        set_font(r2, size=10.7)
    else:
        r = p.add_run(text)
        set_font(r, size=10.7)
    return p


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Trang ")
    set_font(run, size=8.5, color=GRAY)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def add_toc_field(doc):
    p = doc.add_paragraph()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), 'TOC \\o "1-3" \\h \\z \\u')
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "Mục lục sẽ được Word cập nhật tự động khi mở tài liệu."
    r.append(t)
    fld.append(r)
    p._p.append(fld)


def font(size, bold=False):
    path = Path(r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf")
    return ImageFont.truetype(str(path), size)


def draw_box(draw, xy, title, body, fill, outline=NAVY):
    if not outline.startswith("#"):
        outline = f"#{outline}"
    draw.rounded_rectangle(xy, radius=20, fill=fill, outline=outline, width=3)
    x1, y1, x2, y2 = xy
    draw.text((x1 + 22, y1 + 18), title, font=font(28, True), fill="#163A5F")
    y = y1 + 62
    for line in body:
        draw.text((x1 + 25, y), line, font=font(21), fill="#263442")
        y += 32


def arrow(draw, start, end, color="#2E74B5"):
    draw.line([start, end], fill=color, width=6)
    x, y = end
    draw.polygon([(x, y), (x - 16, y - 10), (x - 16, y + 10)], fill=color)


def make_diagrams():
    img = Image.new("RGB", (1800, 980), "white")
    d = ImageDraw.Draw(img)
    d.text((70, 40), "KIẾN TRÚC TỔNG THỂ MUSIC IBOOK", font=font(38, True), fill="#163A5F")
    draw_box(d, (80, 150, 430, 430), "Flutter Client", ["Web / Windows", "Android / iOS", "Provider + GoRouter", "Dio + JWT"], "#E8EEF5")
    draw_box(d, (570, 150, 920, 430), "ASP.NET Core API", ["Controllers", "Business Services", "JWT / Rate limit", "Static audio"], "#F4F8FC")
    draw_box(d, (1060, 150, 1410, 430), "Data & Storage", ["EF Core 8", "SQL Server", "11 bảng nghiệp vụ", "Audio filesystem"], "#FFF7E6")
    draw_box(d, (570, 590, 920, 860), "Music Pipeline", ["MIDI parser", "MusicXML generator", "OSMD SVG", "Canvas practice"], "#EAF7F0")
    draw_box(d, (1060, 590, 1410, 860), "External Runtime", ["Gmail SMTP", "OSMD CDN", "WebView engines", "Secure storage"], "#FDECEC")
    arrow(d, (430, 290), (570, 290)); arrow(d, (920, 290), (1060, 290)); arrow(d, (745, 430), (745, 590)); arrow(d, (920, 720), (1060, 720))
    d.text((1460, 195), "HTTPS/JSON", font=font(21, True), fill="#5B6573")
    arch = ASSETS / "architecture.png"; img.save(arch)

    img = Image.new("RGB", (1800, 1120), "white")
    d = ImageDraw.Draw(img)
    d.text((60, 35), "BẢN ĐỒ MÀN HÌNH VÀ ĐIỀU HƯỚNG", font=font(38, True), fill="#163A5F")
    draw_box(d, (660, 115, 1140, 280), "Public / Identity", ["Login · Register · Forgot OTP · Reset"], "#E8EEF5")
    draw_box(d, (100, 400, 760, 610), "Teacher Shell", ["Tổng quan · Bài học · Học sinh", "Profile"], "#FFF7E6")
    draw_box(d, (1040, 400, 1700, 610), "Student Home", ["Danh sách bài đã xuất bản", "History · Profile"], "#EAF7F0")
    draw_box(d, (100, 770, 760, 1010), "Teacher Workflows", ["Lesson Editor", "Student Detail", "Lesson Analytics", "Assign supplementary work"], "#F4F8FC")
    draw_box(d, (1040, 770, 1700, 1010), "Student Workflows", ["Lesson Detail", "Practice / Section / Exercise", "Exam", "Practice History"], "#F4F8FC")
    arrow(d, (780, 280), (500, 400)); arrow(d, (1020, 280), (1370, 400)); arrow(d, (430, 610), (430, 770)); arrow(d, (1370, 610), (1370, 770))
    sitemap = ASSETS / "screen_map.png"; img.save(sitemap)

    img = Image.new("RGB", (1900, 1250), "white")
    d = ImageDraw.Draw(img)
    d.text((60, 30), "ERD KHÁI QUÁT - 11 BẢNG NGHIỆP VỤ", font=font(38, True), fill="#163A5F")
    boxes = {
        "Users": (760, 100, 1140, 230), "Lessons": (760, 360, 1140, 490),
        "PasswordResetTokens": (80, 100, 540, 230), "LessonNotes": (80, 650, 500, 780),
        "LessonSections": (540, 650, 960, 780), "LessonAnnotations": (1000, 650, 1420, 780),
        "LessonExercises": (1460, 650, 1860, 780), "StudentAssignments": (1400, 980, 1860, 1110),
        "StudentLessonProgresses": (80, 980, 560, 1110), "PracticeSessions": (650, 980, 1050, 1110),
        "StudentNoteAttempts": (650, 1160, 1050, 1235),
    }
    fills = ["#E8EEF5", "#FFF7E6", "#FDECEC", "#F4F8FC", "#F4F8FC", "#F4F8FC", "#F4F8FC", "#EAF7F0", "#EAF7F0", "#EAF7F0", "#FDECEC"]
    for (name, xy), fill in zip(boxes.items(), fills):
        d.rounded_rectangle(xy, radius=18, fill=fill, outline="#2E74B5", width=3)
        d.text((xy[0] + 22, xy[1] + 24), name, font=font(24, True), fill="#163A5F")
    def center(name):
        x1,y1,x2,y2=boxes[name]; return ((x1+x2)//2,(y1+y2)//2)
    # Relationship lines are intentionally schematic to keep the ERD readable.
    for target in ["PasswordResetTokens", "Lessons", "StudentAssignments", "StudentLessonProgresses", "PracticeSessions"]:
        d.line([center("Users"), center(target)], fill="#8292A2", width=3)
    for target in ["LessonNotes", "LessonSections", "LessonAnnotations", "LessonExercises", "StudentAssignments", "StudentLessonProgresses", "PracticeSessions"]:
        d.line([center("Lessons"), center(target)], fill="#2E74B5", width=3)
    d.line([center("PracticeSessions"), center("StudentNoteAttempts")], fill="#B42318", width=3)
    d.line([center("LessonNotes"), center("StudentNoteAttempts")], fill="#B42318", width=3)
    d.text((1280, 1165), "Quy ước: đường nối biểu diễn quan hệ 1-N", font=font(20), fill="#5B6573")
    erd = ASSETS / "erd.png"; img.save(erd)
    return arch, sitemap, erd


features = [
    ("AUTH-01", "Đăng ký tài khoản", "Public", "Đã triển khai", "Tạo Student/Teacher và đăng nhập ngay"),
    ("AUTH-02", "Đăng nhập JWT", "Public", "Đã triển khai", "BCrypt, JWT 7 ngày, secure storage"),
    ("AUTH-03", "Quên mật khẩu OTP Gmail", "Public", "Đã triển khai", "OTP 6 số, hạn 10 phút, rate limit"),
    ("AUTH-04", "Hồ sơ và đổi mật khẩu", "Dùng chung", "Đã triển khai", "Cập nhật tên, mật khẩu tùy chọn, cấp JWT mới"),
    ("LES-01", "Vòng đời bài học", "Teacher", "Đã triển khai", "Tạo nháp, cập nhật, lưu nội dung, xuất bản, xóa"),
    ("MIDI-01", "Import MIDI", "Teacher", "Đã triển khai", "Bảo toàn tempo, nhịp, hóa biểu, track, voice"),
    ("SCORE-01", "MusicXML chuẩn", "Teacher/Student", "Đã triển khai", "Measures, rests, beams, ties, chords; không bịa dữ liệu"),
    ("SCORE-02", "Render OSMD đa nền tảng", "Teacher/Student", "Đã triển khai", "Web iframe, Windows/Android/iOS WebView"),
    ("EDIT-01", "Biên tập nốt", "Teacher", "Đã triển khai", "Pitch, beat, duration, lyric, chord, fingering, staff, voice"),
    ("EDIT-02", "Công cụ soạn nhạc", "Teacher", "Đã triển khai", "Transport, metronome, quantize, undo/redo, timeline, piano"),
    ("PED-01", "Đoạn luyện tập", "Teacher/Student", "Đã triển khai", "Beat range, độ khó, BPM, tay trái/phải/cả hai"),
    ("PED-02", "Ghi chú sư phạm", "Teacher/Student", "Đã triển khai", "Ngón, sắc thái, legato, pedal, tempo, ghi chú"),
    ("PED-03", "Thiết kế bài tập", "Teacher/Student", "Đã triển khai", "6 loại bài tập và cấu hình JSON"),
    ("LEARN-01", "Luyện tập thời gian thực", "Student", "Đã triển khai", "Countdown, timeline, piano, phán định và MISS tự động"),
    ("LEARN-02", "Kiểm tra", "Student", "Đã triển khai", "Cửa sổ timing chặt, báo cáo lỗi chi tiết"),
    ("LEARN-03", "Bài tập bổ sung", "Teacher/Student", "Đã triển khai", "Giao theo bài/đoạn/exercise và tự hoàn thành sau luyện"),
    ("ASSESS-01", "Chấm từng nốt", "Student/Teacher", "Đã triển khai", "Pitch, timing, judge result, score, accuracy"),
    ("ASSESS-02", "Lịch sử và tiến bộ", "Student/Teacher", "Đã triển khai", "Session history, biểu đồ, nốt sai, thời gian UTC+7"),
    ("AUDIO-01", "Audio tham chiếu", "Teacher/Student", "API có, UI chưa nối", "Upload endpoint tồn tại; editor chưa kết nối"),
    ("PROG-01", "Trang tiến độ riêng học sinh", "Student", "API có, UI placeholder", "Route chưa hoạt động"),
]

screens = [
    ("SCR-01", "/login", "Public", "Đăng nhập", "Email/mật khẩu, validate, hiện/ẩn mật khẩu", "POST /api/auth/login"),
    ("SCR-02", "/register", "Public", "Đăng ký", "Họ tên, email, mật khẩu, chọn vai trò", "POST /api/auth/register"),
    ("SCR-03", "/forgot-password", "Public", "Yêu cầu OTP", "Nhập email, trạng thái gửi/lỗi/thành công", "POST /api/auth/forgot-password"),
    ("SCR-04", "/reset-password?email=", "Public", "Đặt lại mật khẩu", "OTP 6 số, mật khẩu mới >=8 ký tự", "POST /api/auth/reset-password"),
    ("SCR-05", "/profile", "Teacher/Student", "Hồ sơ", "Avatar chữ, tên, vai trò, đổi mật khẩu, logout", "PUT /api/auth/profile"),
    ("SCR-06", "/teacher", "Teacher", "Dashboard quản trị", "3 tab Tổng quan/Bài học/Học sinh", "Dashboard + lessons/mine + progress"),
    ("SCR-07", "/teacher/lesson/new", "Teacher", "Tạo bài", "Editor chuyên sâu, import MIDI và lưu nháp", "POST lesson + content"),
    ("SCR-08", "/teacher/lesson/{id}", "Teacher", "Soạn/sửa bài", "Khuông, timeline, piano, section, annotation, exercise", "GET/PUT/content/publish"),
    ("SCR-09", "/teacher/student/{id}", "Teacher", "Chi tiết học sinh", "Tiến độ theo bài, lần luyện/kiểm tra gần nhất UTC+7", "GET student progress"),
    ("SCR-10", "/teacher/student/{studentId}/lesson/{lessonId}/analytics", "Teacher", "Phân tích bài", "Biểu đồ, nốt sai, session, giao bổ sung", "GET analytics; POST assignment"),
    ("SCR-11", "/student", "Student", "Trang bài học", "Danh sách bài xuất bản, history/profile, grid responsive", "GET /api/lessons"),
    ("SCR-12", "/student/lesson/{id}", "Student", "Chi tiết bài", "Khuông, đoạn luyện, exercise, assignment, ghi chú", "GET lesson + assignments"),
    ("SCR-13", "/student/practice/{id}", "Student", "Luyện tập", "Toàn bài/đoạn/exercise/assignment, realtime judge", "GET lesson; POST practice"),
    ("SCR-14", "/student/exam/{id}", "Student", "Kiểm tra", "Timing ±300ms, báo cáo tổng quan/chi tiết lỗi", "GET lesson; POST practice"),
    ("SCR-15", "/student/history", "Student", "Lịch sử", "Session, score, accuracy, chi tiết từng nốt", "GET practice-history"),
]

api_rows = [
    ("POST", "/api/auth/register", "Public", "RegisterRequest", "AuthResponse", "Đăng ký"),
    ("POST", "/api/auth/login", "Public", "LoginRequest", "AuthResponse", "Đăng nhập"),
    ("POST", "/api/auth/forgot-password", "Public + 5/IP/phút", "email", "message", "Gửi OTP"),
    ("POST", "/api/auth/reset-password", "Public", "email, otp, newPassword", "message", "Đặt lại mật khẩu"),
    ("PUT", "/api/auth/profile", "JWT", "fullName, newPassword?", "AuthResponse", "Cập nhật hồ sơ"),
    ("GET", "/api/lessons", "Public", "-", "Published lessons", "Danh sách học sinh"),
    ("GET", "/api/lessons/mine", "Teacher", "-", "Teacher lessons", "Quản lý bài"),
    ("GET", "/api/lessons/{id}", "Public hiện tại", "id", "Lesson detail", "Mở bài/editor"),
    ("POST", "/api/lessons", "Teacher", "Lesson metadata", "Lesson", "Tạo nháp"),
    ("PUT", "/api/lessons/{id}", "Teacher", "Lesson metadata", "Lesson", "Sửa metadata"),
    ("DELETE", "/api/lessons/{id}", "Teacher", "id", "message", "Xóa bài + phụ thuộc"),
    ("POST", "/api/lessons/{id}/content", "Teacher", "notes/sections/annotations/exercises", "200", "Lưu nội dung atomic"),
    ("POST", "/api/lessons/{id}/notes", "Teacher", "LessonNote", "LessonNote", "Thêm nốt đơn"),
    ("DELETE", "/api/lessons/notes/{noteId}", "Teacher", "noteId", "message", "Xóa nốt"),
    ("DELETE", "/api/lessons/{id}/notes", "Teacher", "id", "200", "Xóa toàn bộ nốt"),
    ("POST", "/api/lessons/{id}/audio", "Teacher", "multipart file", "audioUrl", "Upload audio"),
    ("PUT", "/api/lessons/{id}/publish", "Teacher", "id", "message", "Xuất bản"),
    ("GET", "/api/student/progress", "Student", "-", "Progress[]", "Tiến độ cá nhân"),
    ("PUT", "/api/student/progress/{lessonId}", "Student", "progress payload", "Progress", "Lưu tiến độ"),
    ("POST", "/api/student/practice", "Student", "session + attempts[]", "PracticeSession", "Chấm/lưu phiên"),
    ("GET", "/api/student/practice-history", "Student", "-", "Sessions + attempts", "Lịch sử"),
    ("GET", "/api/student/assignments", "Student", "lessonId?", "AssignmentDto[]", "Bài bổ sung"),
    ("GET", "/api/teacher/dashboard", "Teacher", "-", "DashboardDto", "Chỉ số tổng quan"),
    ("GET", "/api/teacher/students-progress", "Teacher", "-", "StudentProgress[]", "Danh sách học sinh"),
    ("GET", "/api/teacher/students/{id}/progress", "Teacher", "studentId", "Student detail", "Tiến độ chi tiết"),
    ("GET", "/api/teacher/students/{sid}/lessons/{lid}/analytics", "Teacher", "studentId, lessonId", "AnalyticsDto", "Phân tích"),
    ("POST", "/api/teacher/assignments", "Teacher", "Assignment request", "AssignmentDto", "Giao bổ sung"),
]

tables = [
    ("Users", "Tài khoản, phân quyền và xác thực", "Login, Register, Profile, Teacher student management", [
        ("Id", "bigint identity", "PK, not null", "Định danh người dùng"), ("FullName", "nvarchar(max)", "not null", "Họ tên hiển thị"),
        ("Email", "nvarchar(450)", "unique, not null", "Email đăng nhập"), ("PasswordHash", "nvarchar(max)", "nullable", "BCrypt hash; nullable cho legacy Google"),
        ("AuthProvider", "nvarchar(max)", "not null", "Nguồn xác thực, hiện dùng Local"), ("ProviderKey", "nvarchar(max)", "nullable", "ID nhà cung cấp cũ"),
        ("Role", "nvarchar(max)", "not null", "Teacher hoặc Student"), ("CreatedAtUtc", "datetime2", "not null", "Thời điểm tạo UTC")]),
    ("PasswordResetTokens", "OTP/token đặt lại mật khẩu", "Forgot Password, Reset Password", [
        ("Id", "bigint identity", "PK", "Định danh token"), ("UserId", "bigint", "FK Users", "Chủ sở hữu yêu cầu"),
        ("Token", "nvarchar(max)", "not null", "HMAC của OTP, không lưu OTP rõ"), ("ExpiresAtUtc", "datetime2", "not null", "Hết hạn sau 10 phút"),
        ("IsUsed", "bit", "not null", "Ngăn sử dụng lại")]),
    ("Lessons", "Bản ghi gốc của bài học/bản nhạc", "Teacher editor/list; Student home/detail/practice/exam", [
        ("Id", "bigint identity", "PK", "Định danh bài"), ("TeacherId", "bigint", "FK Users, Restrict", "Giáo viên sở hữu"),
        ("Title", "nvarchar(max)", "not null", "Tên bài"), ("Composer", "nvarchar(max)", "not null", "Tác giả"),
        ("Clef", "nvarchar(max)", "not null", "Khóa nhạc mặc định"), ("KeySignature", "nvarchar(max)", "not null", "Hóa biểu gốc"),
        ("TimeSignature", "nvarchar(max)", "not null", "Số chỉ nhịp chính"), ("TimeSignatureMap", "nvarchar(max)", "not null", "Các thay đổi nhịp theo MIDI"),
        ("TempoMap", "nvarchar(max)", "not null", "Các thay đổi tempo"), ("Tempo", "int", "not null", "BPM mặc định/gốc"),
        ("AudioUrl", "nvarchar(max)", "nullable", "URL audio tham chiếu"), ("AudioFileName", "nvarchar(max)", "nullable", "Tên file upload"),
        ("CreatedAtUtc", "datetime2", "not null", "Ngày tạo UTC"), ("IsPublished", "bit", "not null", "Nháp/đã xuất bản")]),
    ("LessonNotes", "Dữ liệu từng nốt dùng chung để render, phát và chấm", "Teacher editor; Student detail/practice/exam; analytics", [
        ("Id", "bigint identity", "PK", "Định danh nốt"), ("LessonId", "bigint", "FK Lessons, Cascade", "Bài chứa nốt"),
        ("Second", "float", "not null", "Thời điểm bắt đầu theo giây"), ("StartBeat", "float", "not null", "Vị trí theo phách"),
        ("DurationBeat", "float", "not null", "Trường độ theo phách"), ("Velocity", "int", "not null", "Cường độ MIDI"),
        ("Track", "int", "not null", "Số track MIDI"), ("TrackName", "nvarchar(max)", "not null", "Tên track MIDI"),
        ("Staff", "int", "not null", "Khuông"), ("Voice", "int", "not null", "Voice độc lập"),
        ("Note", "nvarchar(max)", "not null", "Cao độ C4/F#4..."), ("Duration", "nvarchar(max)", "not null", "Tên trường độ ký âm"),
        ("Lyric", "nvarchar(max)", "not null", "Lời chỉ khi có/được nhập"), ("Chord", "nvarchar(max)", "not null", "Ký hiệu hợp âm"),
        ("Fingering", "nvarchar(max)", "not null", "Số ngón đàn")]),
    ("LessonSections", "Chia bài thành đoạn luyện tập", "Teacher authoring; Student lesson/practice", [
        ("Id", "bigint identity", "PK", "Định danh đoạn"), ("LessonId", "bigint", "FK Lessons, Cascade", "Bài cha"),
        ("Title", "nvarchar(max)", "not null", "Tên đoạn"), ("StartBeat", "float", "not null", "Phách bắt đầu"),
        ("EndBeat", "float", "not null", "Phách kết thúc"), ("DefaultTempo", "int", "30-300", "BPM của đoạn"),
        ("Difficulty", "nvarchar(max)", "not null", "Beginner/Intermediate/Advanced"), ("Hand", "nvarchar(max)", "not null", "Both/Right/Left"),
        ("SortOrder", "int", "unique per lesson", "Thứ tự hiển thị")]),
    ("LessonAnnotations", "Ghi chú/ký hiệu sư phạm ngoài MIDI", "Teacher authoring; Student lesson detail", [
        ("Id", "bigint identity", "PK", "Định danh ghi chú"), ("LessonId", "bigint", "FK Lessons, Cascade", "Bài chứa ghi chú"),
        ("StartBeat", "float", "not null", "Phách bắt đầu"), ("EndBeat", "float", "nullable", "Phách kết thúc"),
        ("Kind", "nvarchar(max)", "not null", "Finger/Dynamic/Articulation/Pedal/Tempo/TeacherNote"), ("Text", "nvarchar(max)", "not null", "Nội dung hiển thị")]),
    ("LessonExercises", "Cấu hình bài tập theo toàn bài hoặc section", "Teacher authoring; Student lesson/practice", [
        ("Id", "bigint identity", "PK", "Định danh exercise"), ("LessonId", "bigint", "FK Lessons", "Bài cha"),
        ("LessonSectionId", "bigint", "nullable, FK Section", "Null = toàn bài; SetNull khi xóa section"), ("Title", "nvarchar(max)", "not null", "Tên bài tập"),
        ("Type", "nvarchar(max)", "not null", "6 loại bài tập"), ("Instruction", "nvarchar(max)", "not null", "Hướng dẫn học sinh"),
        ("ConfigJson", "nvarchar(max)", "not null", "Cấu hình riêng theo loại"), ("SortOrder", "int", "indexed", "Thứ tự hiển thị")]),
    ("StudentAssignments", "Bài bổ sung giao riêng cho học sinh", "Teacher analytics/assign; Student lesson detail/practice", [
        ("Id", "bigint identity", "PK", "Định danh giao bài"), ("StudentId", "bigint", "FK Users, Restrict", "Học sinh nhận"),
        ("LessonId", "bigint", "FK Lessons, Restrict", "Bài được giao"), ("LessonSectionId", "bigint", "nullable, SetNull", "Đoạn cụ thể"),
        ("LessonExerciseId", "bigint", "nullable, SetNull", "Exercise cụ thể"), ("Message", "nvarchar(max)", "not null", "Lời dặn"),
        ("DueAtUtc", "datetime2", "nullable", "Hạn hoàn thành"), ("CreatedAtUtc", "datetime2", "not null", "Ngày giao"),
        ("IsCompleted", "bit", "not null", "Trạng thái hoàn thành")]),
    ("StudentLessonProgresses", "Tiến độ tổng hợp mỗi học sinh-bài", "Student detail/progress; Teacher student management", [
        ("Id", "bigint identity", "PK", "Định danh"), ("StudentId", "bigint", "FK Users", "Học sinh"),
        ("LessonId", "bigint", "FK Lessons", "Bài học"), ("LastPositionSecond", "float", "not null", "Vị trí gần nhất"),
        ("CompletedNoteCount", "int", "not null", "Số nốt hoàn thành"), ("BestScore", "int", "not null", "Điểm cao nhất"),
        ("TotalAttempts", "int", "not null", "Tổng lần luyện/thi"), ("IsCompleted", "bit", "not null", "Đã hoàn thành"),
        ("LastStudiedAtUtc", "datetime2", "not null", "Lần học gần nhất UTC")]),
    ("PracticeSessions", "Mỗi phiên luyện tập hoặc kiểm tra", "Practice/Exam result; History; Teacher analytics", [
        ("Id", "bigint identity", "PK", "Định danh phiên"), ("StudentId", "bigint", "FK Users", "Học sinh"),
        ("LessonId", "bigint", "FK Lessons", "Bài luyện"), ("IsExam", "bit", "not null", "Luyện hay kiểm tra"),
        ("Score", "int", "not null", "Điểm phiên"), ("CorrectCount", "int", "not null", "Số nốt đúng"),
        ("WrongCount", "int", "not null", "Số nốt sai"), ("Accuracy", "float", "not null", "Tỉ lệ chính xác"),
        ("DurationSeconds", "int", "not null", "Thời lượng"), ("StartedAtUtc", "datetime2", "not null", "Bắt đầu UTC"),
        ("FinishedAtUtc", "datetime2", "nullable", "Kết thúc UTC"), ("CreatedAtUtc", "datetime2", "not null", "Ngày tạo bản ghi")]),
    ("StudentNoteAttempts", "Kết quả chi tiết từng nốt trong phiên", "Practice/Exam report; History; Teacher error analytics", [
        ("Id", "bigint identity", "PK", "Định danh attempt"), ("PracticeSessionId", "bigint", "FK Session, Cascade", "Phiên cha"),
        ("LessonNoteId", "bigint", "FK Note, Restrict", "Nốt chuẩn"), ("ExpectedNote", "nvarchar(max)", "not null", "Nốt cần chơi"),
        ("PlayedNote", "nvarchar(max)", "not null", "Nốt đã chơi"), ("ExpectedAtSecond", "float", "not null", "Thời điểm chuẩn"),
        ("PlayedAtSecond", "float", "not null", "Thời điểm thực"), ("TimingErrorMs", "float", "not null", "Sai lệch ms"),
        ("IsCorrectPitch", "bit", "not null", "Đúng cao độ"), ("IsCorrectTiming", "bit", "not null", "Đúng nhịp"),
        ("IsCorrect", "bit", "not null", "Kết quả chung"), ("JudgeResult", "nvarchar(max)", "not null", "PERFECT/GOOD/LATE/WRONG/MISS"),
        ("CreatedAtUtc", "datetime2", "not null", "Thời điểm ghi nhận")]),
]


def build_document():
    arch, sitemap, erd = make_diagrams()
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5); sec.page_height = Inches(11)
    sec.top_margin = Inches(0.82); sec.bottom_margin = Inches(0.78)
    sec.left_margin = Inches(1); sec.right_margin = Inches(1)
    sec.header_distance = Inches(0.35); sec.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"; normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.25
    for level, size, color, before, after in [(1,16,BLUE,18,10),(2,13,BLUE,14,7),(3,12,NAVY,10,5)]:
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Calibri"; style.font.size = Pt(size); style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before); style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    # Running header/footer.
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hr = hp.add_run("MUSIC IBOOK  |  TÀI LIỆU QUẢN LÝ DỰ ÁN")
    set_font(hr, size=8.5, bold=True, color=GRAY)
    add_page_number(sec.footer.paragraphs[0])

    # Editorial cover.
    for _ in range(4): doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("PROJECT REFERENCE GUIDE"); set_font(r, size=10, bold=True, color=GOLD)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(10)
    r = p.add_run("TÀI LIỆU QUẢN LÝ DỰ ÁN"); set_font(r, size=28, bold=True, color=NAVY)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(8)
    r = p.add_run("MUSIC IBOOK"); set_font(r, size=24, bold=True, color=BLUE)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(34)
    r = p.add_run("Nền tảng biên soạn, luyện tập và đánh giá piano đa nền tảng"); set_font(r, size=13, italic=True, color=GRAY)
    add_callout(doc, "Phạm vi tài liệu", "Thiết kế database · thiết kế màn hình · screen list · feature list · API list · luồng nghiệp vụ · ma trận truy vết · bảo mật · kiểm thử · rủi ro và lộ trình.", fill=PALE_BLUE)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(40)
    r = p.add_run("Phiên bản 1.0  |  Cơ sở: mã nguồn hiện hành"); set_font(r, size=10.5, bold=True, color=NAVY)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Phân loại: Tài liệu nội bộ dự án"); set_font(r, size=9.5, color=GRAY)
    doc.add_page_break()

    add_heading(doc, "Kiểm soát tài liệu", 1)
    add_table(doc, ["Thuộc tính", "Giá trị"], [
        ("Tên dự án", "Music IBook"), ("Loại tài liệu", "Project Management & System Design Reference"),
        ("Phiên bản", "1.0"), ("Trạng thái", "Baseline theo mã nguồn hiện tại"),
        ("Đối tượng đọc", "Product Owner, Project Manager, BA, UI/UX, Flutter, Backend, QA, DevOps"),
        ("Nguồn xác minh", "Flutter source, ASP.NET Core source, EF Core models/migrations/snapshot"),
        ("Nguyên tắc bảo mật", "Không ghi secret, app password Gmail, JWT key hoặc connection string"),
    ], [2700, 6660], font_size=9.2)
    add_heading(doc, "Quy ước trạng thái", 2)
    add_table(doc, ["Trạng thái", "Ý nghĩa"], [
        ("Đã triển khai", "Có code và đã nối vào luồng UI/API hiện hành."),
        ("API có, UI chưa nối", "Backend/service đã có nhưng người dùng chưa truy cập được từ màn hình chính."),
        ("Placeholder", "Tệp hoặc ý tưởng tồn tại nhưng chưa có triển khai sử dụng được."),
        ("Đề xuất", "Thiết kế mục tiêu hoặc cải tiến, chưa phải hành vi hiện tại."),
    ], [2200, 7160], font_size=9.2)
    add_heading(doc, "Mục lục", 1); add_toc_field(doc)
    add_body(doc, "Lưu ý: trong Microsoft Word, nhấn Ctrl+A rồi F9 để cập nhật mục lục và số trang nếu cần.")
    doc.add_page_break()

    add_heading(doc, "1. Tổng quan dự án", 1)
    add_body(doc, "Music IBook là ứng dụng học piano đa nền tảng kết nối quy trình biên soạn của giáo viên với luyện tập, kiểm tra và phân tích tiến bộ của học sinh. Giáo viên có thể nhập MIDI, hiệu chỉnh bản nhạc, chia đoạn, thiết kế bài tập và xuất bản. Học sinh xem cùng một bản nhạc, luyện theo đoạn hoặc bài giao, nhận phản hồi thời gian thực và lưu lịch sử chi tiết từng nốt.")
    add_heading(doc, "1.1 Mục tiêu sản phẩm", 2)
    for text in ["Chuẩn hóa MIDI thành bản nhạc có measures, rests, beams, ties và giữ metadata gốc.", "Tạo một nguồn dữ liệu khuông nhạc dùng chung để teacher và student hiển thị nhất quán.", "Đưa tri thức sư phạm vào bài học qua section, annotation, exercise và assignment.", "Đánh giá pitch/timing ở cấp nốt và biến dữ liệu luyện tập thành thông tin hành động cho giáo viên.", "Đảm bảo trải nghiệm Web, Windows và mobile có khả năng cuộn, responsive và hiệu năng phù hợp."]:
        add_bullet(doc, text)
    add_heading(doc, "1.2 Phạm vi hiện hành", 2)
    add_table(doc, ["Trong phạm vi", "Ngoài/chưa hoàn chỉnh"], [
        ("Tài khoản local, JWT, OTP Gmail", "Google Login đã loại bỏ; cột legacy còn trong DB"),
        ("MIDI parse tại Flutter và lưu note metadata", "Backend chưa nhận file MIDI/không trả MusicXML"),
        ("OSMD read-only + canvas tương tác", "Phụ thuộc CDN OSMD khi chạy online"),
        ("Practice/exam/history/teacher analytics", "Chưa có mô hình lớp hoặc quan hệ teacher-student"),
        ("Audio upload API", "Màn hình upload/editor chưa nối hoàn chỉnh"),
    ], [4680, 4680], font_size=9)
    add_heading(doc, "1.3 Vai trò và bên liên quan", 2)
    add_table(doc, ["Vai trò", "Trách nhiệm/nhu cầu"], [
        ("Giáo viên", "Tạo và xuất bản bài, thiết kế nội dung sư phạm, theo dõi tiến bộ, giao bổ sung."),
        ("Học sinh", "Khám phá bài, luyện/thi, nhận phản hồi, xem lịch sử và hoàn thành assignment."),
        ("Product Owner/PM", "Ưu tiên backlog, quản lý phạm vi, tiêu chí nghiệm thu và release."),
        ("Flutter team", "UI đa nền tảng, MIDI/MusicXML, OSMD/WebView, realtime practice."),
        ("Backend team", "API, auth, chấm/lưu session, analytics, database và security."),
        ("QA/DevOps", "Test cross-platform, migration, secret, deployment, monitoring và rollback."),
    ], [2200, 7160], font_size=9.1)

    add_heading(doc, "2. Kiến trúc hệ thống", 1)
    doc.add_picture(str(arch), width=Inches(6.45)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph("Hình 1. Kiến trúc logic tổng thể"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; set_font(p.runs[0], size=9, italic=True, color=GRAY)
    add_heading(doc, "2.1 Kiến trúc client", 2)
    add_body(doc, "Flutter sử dụng MaterialApp.router và go_router. Mỗi màn hình tương tác với ChangeNotifier Provider, Provider gọi Service, Service sử dụng singleton Dio ApiClient. JWT được lấy từ flutter_secure_storage và gắn tự động qua interceptor.")
    add_table(doc, ["Thành phần", "Công nghệ", "Trách nhiệm"], [
        ("Routing", "go_router", "Route guard đăng nhập và điều hướng theo vai trò sau auth"),
        ("State", "provider/ChangeNotifier", "Auth, lesson, student, history, progress"),
        ("HTTP", "Dio", "Base URL theo platform, JWT interceptor, timeout/error mapping"),
        ("Notation", "MusicXmlGenerator + OSMD", "Sinh MusicXML và render SVG read-only"),
        ("Realtime score", "Custom canvas engraving", "Hit geometry, highlight, cuộn theo playhead ở 60fps"),
        ("Persistence", "flutter_secure_storage", "Token, role và tên người dùng"),
    ], [1900, 2300, 5160], font_size=9)
    add_heading(doc, "2.2 Kiến trúc backend", 2)
    add_body(doc, "ASP.NET Core 8 triển khai chuỗi Controller → Service → EF Core DbContext/Repository → SQL Server. Auth dùng generic repository; Lesson, Student và Teacher truy cập DbContext trong service. API tự chạy migration và seed khi khởi động.")
    add_heading(doc, "2.3 Pipeline bản nhạc", 2)
    for step in ["Chọn file MIDI tại Flutter và parse trong isolate compute.", "Đọc format 0/1, tempo map, time-signature map, key signature; bỏ percussion channel 10.", "Giữ track, trackName, staff, voice; quantize beat và trường độ; không tự tạo lyric/annotation.", "Lưu lesson metadata và content nguyên khối qua API.", "Sinh MusicXML theo part/track, measures, rests, beams, ties, chord shared stem và metadata có thật.", "Render read-only bằng OSMD; practice/exam dùng canvas tương tác nhưng cùng nguồn LessonNote."]:
        add_number(doc, step)
    add_callout(doc, "Nguyên tắc nguồn dữ liệu duy nhất", "Teacher và student đều đọc cùng MusicLesson/LessonNote và dùng cùng MusicXmlGenerator. Đây là điều kiện cốt lõi để khuông nhạc hai phía giống nhau.")

    add_heading(doc, "3. Danh mục tính năng", 1)
    add_table(doc, ["ID", "Tính năng", "Vai trò", "Trạng thái", "Ghi chú"], features, [1050, 2100, 1350, 1600, 3260], font_size=8.4)
    add_heading(doc, "3.1 Quy tắc nghiệp vụ chính", 2)
    for rule in ["Bài mới luôn là nháp; chỉ bài IsPublished=true xuất hiện trong danh sách học sinh.", "Section phải có title, StartBeat >=1, EndBeat >= StartBeat, BPM 30–300 và SortOrder duy nhất trong bài.", "Exercise có thể áp dụng toàn bài hoặc gắn section qua sectionSortOrder khi lưu transaction.", "OTP 6 số có hiệu lực 10 phút; OTP mới vô hiệu các token cũ chưa dùng.", "Practice: cửa sổ ±500ms, đúng mới tiến nốt; Exam: ±300ms và luôn tiến nốt sau input.", "JudgeResult hiện gồm PERFECT, GOOD, LATE, WRONG, MISS; điểm tương ứng 10, 8, 5, -5, -10 và tổng không âm.", "Ngày giờ lưu UTC; giao diện quản lý hiển thị cố định UTC+7 bằng VietnamTime."]:
        add_bullet(doc, rule)

    add_heading(doc, "4. Thiết kế màn hình", 1)
    doc.add_picture(str(sitemap), width=Inches(6.45)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph("Hình 2. Bản đồ điều hướng màn hình đang hoạt động"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; set_font(p.runs[0], size=9, italic=True, color=GRAY)
    add_heading(doc, "4.1 Nguyên tắc UI/UX", 2)
    for text in ["Ưu tiên khuông nhạc và playhead là vùng nội dung chính; công cụ sư phạm đặt theo ngữ cảnh.", "TextField phải chặn hotkey piano để nhập metadata không phát nốt ngoài ý muốn.", "Loading, empty, retry và error phải là trạng thái thiết kế chính thức, không hiển thị stack trace.", "Teacher editor trên mobile cuộn dọc; SVG rộng 100%, ẩn cuộn ngang và touch-action pan-y.", "Các thao tác xóa/xuất bản cần xác nhận và phản hồi SnackBar rõ ràng.", "Màu sắc phản hồi PERFECT/GOOD/LATE/WRONG/MISS nhất quán trong practice, exam và history."]:
        add_bullet(doc, text)
    add_heading(doc, "4.2 Screen list", 2)
    add_table(doc, ["ID", "Route", "Role", "Mục đích", "Chức năng chính", "API"], screens, [720, 1820, 1050, 1450, 2500, 1820], font_size=7.7)
    add_heading(doc, "4.3 Responsive và platform", 2)
    add_table(doc, ["Khu vực", "Mobile", "Tablet/Desktop"], [
        ("Student Home", "<600: 1 cột", ">=600: 2 cột; >=1024: 3 cột"),
        ("Lesson Detail", "<900: 2 tab; header/dock rút gọn khi <600 hoặc cao <500", ">=900: split 40/60"),
        ("Teacher Editor", "Khuông cao max(360, 48% màn hình), timeline 72, cuộn dọc", "Khuông 460, timeline 92, thanh công cụ đầy đủ"),
        ("OSMD", "SVG width 100%, overflow-y auto, pan-y, cuộn quán tính", "Endless page, vertical scroll, zoom"),
        ("Renderer", "webview_flutter", "Web iframe; Windows webview_windows"),
    ], [1800, 3780, 3780], font_size=8.7)

    details = [
        ("4.4 Nhóm Identity", "SCR-01 đến SCR-05", ["Form 1 cột, label rõ, hỗ trợ hiện/ẩn password và validation tại chỗ.", "Không còn nút Google Login.", "Forgot → Reset truyền email qua query; reset yêu cầu OTP 6 số và password >=8.", "Profile hiển thị avatar chữ cái, role, đổi tên/password và logout."]),
        ("4.5 Teacher Dashboard", "SCR-06", ["AppBar + profile, BottomNavigation 3 tab.", "Dashboard: 5 KPI tổng quan.", "Bài học: list, trạng thái loading/error/retry/empty, edit/delete, FAB Tạo bài.", "Học sinh: tên/email/số bài/số kiểm tra/điểm TB/hoạt động UTC+7."]),
        ("4.6 Teacher Lesson Editor", "SCR-07/SCR-08", ["Header: tên bài, tác giả, tempo, nhịp, hóa biểu; action import/lưu/xuất bản.", "Toolbar: play/pause/stop/record, undo/redo, duplicate/delete, duration, dotted/triplet, quantize, zoom, octave.", "Workspace: lời/hợp âm nốt kế, OSMD score, beat timeline, inspector, piano ảo.", "Pedagogy: section, annotation và exercise editor.", "Focus guard: khóa piano/hotkey khi đang nhập TextField."]),
        ("4.7 Teacher Progress & Analytics", "SCR-09/SCR-10", ["Student detail dùng ExpansionTile theo lesson với trạng thái, score, accuracy và thời gian UTC+7.", "Analytics tải song song lesson và analytics; biểu đồ accuracy, session gần nhất, nốt sai.", "FAB giao bài theo toàn bài/section/exercise; nên bổ sung field due date trong dialog."]),
        ("4.8 Student Learning", "SCR-11 đến SCR-15", ["Home hiển thị lesson published bằng responsive grid.", "Lesson detail tải lesson và assignments song song; hiển thị section/exercise/annotation và dock Practice/Exam.", "Practice hỗ trợ query sectionId/exerciseId/assignmentId; lọc beat/hand/tempo.", "Exam có timing chặt và báo cáo lỗi; History dùng ExpansionTile theo session."]),
    ]
    for heading, ids, bullets in details:
        add_heading(doc, heading, 2); add_body(doc, f"Phạm vi: {ids}", bold_prefix="Phạm vi:")
        for b in bullets: add_bullet(doc, b)
    add_callout(doc, "Placeholder/dead files", "splash_screen.dart, student/progress_screen.dart, theory_screen.dart, upload_audio_screen.dart, lesson_music_editor.dart và lesson_theory_editor.dart hiện rỗng/không có route. Cần xóa hoặc hoàn thiện trước release để giảm nhiễu kiến trúc.", fill="FFF7E6", accent=GOLD)

    add_heading(doc, "5. Thiết kế API", 1)
    add_body(doc, "API sử dụng JSON camelCase, JWT Bearer cho endpoint bảo vệ và phân vai Teacher/Student. Tổng kiểm kê theo controller hiện tại là 27 route vận hành (bao gồm các thao tác note/audio/content chi tiết); nhóm API cốt lõi được liệt kê dưới đây.")
    add_table(doc, ["Method", "Endpoint", "Quyền", "Request", "Response", "Mục đích"], api_rows, [850, 2620, 1400, 1700, 1300, 1490], font_size=7.7)
    add_heading(doc, "5.1 Contract lưu nội dung bài học", 2)
    add_table(doc, ["Mảng", "Trường chính", "Vai trò"], [
        ("notes[]", "second, beat, duration, velocity, track/name, staff, voice, note, lyric, chord, fingering", "Toàn bộ dữ liệu khuông và chấm nốt"),
        ("sections[]", "title, start/end beat, tempo, difficulty, hand, sortOrder", "Đoạn luyện tập"),
        ("annotations[]", "beat range, kind, text", "Ghi chú sư phạm"),
        ("exercises[]", "title, type, instruction, configJson, sectionSortOrder, sortOrder", "Bài tập có cấu hình"),
    ], [1450, 4900, 3010], font_size=8.8)
    add_heading(doc, "5.2 Chuẩn lỗi đề xuất", 2)
    add_body(doc, "Hiện exception chưa được chuẩn hóa đầy đủ. Mục tiêu production nên dùng ProblemDetails thống nhất và không trả stack trace cho client.")
    add_table(doc, ["HTTP", "Trường hợp", "Ví dụ message"], [
        ("400", "Validation sai", "Dữ liệu đoạn luyện tập không hợp lệ"), ("401", "Thiếu/hết hạn JWT", "Authentication required"),
        ("403", "Sai role/không sở hữu", "Bạn không có quyền thao tác bài học này"), ("404", "Không tìm thấy", "Không tìm thấy bài học"),
        ("409", "Xung đột dữ liệu", "SortOrder đã tồn tại"), ("429", "Rate limit", "Vui lòng thử lại sau"),
        ("500", "Lỗi nội bộ", "Mã sự cố; không chứa stack trace"),
    ], [1100, 3300, 4960], font_size=8.8)

    add_heading(doc, "6. Thiết kế database", 1)
    doc.add_picture(str(erd), width=Inches(6.45)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph("Hình 3. ERD khái quát"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; set_font(p.runs[0], size=9, italic=True, color=GRAY)
    add_body(doc, "Nền tảng dùng SQL Server và EF Core 8. Khóa chính nghiệp vụ dùng bigint IDENTITY. Thời gian lưu UTC bằng datetime2 và được chuyển UTC+7 khi hiển thị. Schema hiện có 11 bảng nghiệp vụ.")
    add_heading(doc, "6.1 Quan hệ và chính sách xóa", 2)
    add_table(doc, ["Quan hệ", "Bội số", "Delete behavior/ghi chú"], [
        ("Users(Teacher) → Lessons", "1-N", "Restrict"), ("Lessons → Notes/Sections/Annotations", "1-N", "Cascade"),
        ("Lessons → Exercises", "1-N", "Runtime Restrict; snapshot đang Cascade"), ("Sections → Exercises", "1-N", "SetNull"),
        ("Users/Lessons → Assignments", "1-N", "Restrict"), ("Sections/Exercises → Assignments", "1-N", "SetNull"),
        ("Sessions → NoteAttempts", "1-N", "Cascade"), ("LessonNotes → NoteAttempts", "1-N", "Restrict"),
    ], [3950, 1200, 4210], font_size=8.8)
    add_callout(doc, "Schema drift cần xử lý", "MusicIBookDbContext cấu hình Lesson → LessonExercise là Restrict nhưng ModelSnapshot hiện ghi Cascade. Trước triển khai production phải tạo migration đồng bộ và kiểm thử xóa bài trên database thật.", fill="FDECEC", accent=RED)

    for idx, (name, purpose, use, columns) in enumerate(tables, 1):
        add_heading(doc, f"6.{idx + 1} Bảng {name}", 2)
        add_table(doc, ["Thuộc tính", "Mô tả"], [("Mục đích", purpose), ("Màn hình/luồng phục vụ", use)], [2300, 7060], font_size=9)
        add_table(doc, ["Cột", "Kiểu", "Ràng buộc", "Ý nghĩa"], columns, [1950, 1900, 2100, 3410], font_size=8.25)

    add_heading(doc, "7. Luồng nghiệp vụ end-to-end", 1)
    workflows = [
        ("7.1 Authentication", ["Nhập và validate thông tin.", "Gọi Auth API; backend xác thực BCrypt hoặc xử lý OTP.", "Nhận JWT/role/name và lưu secure storage.", "Router điều hướng Teacher hoặc Student; interceptor gắn JWT cho request sau."]),
        ("7.2 Soạn và xuất bản", ["Import MIDI hoặc nhập bằng piano/editor.", "Chỉnh metadata/nốt; thêm section, annotation, exercise.", "POST/PUT lesson để có ID.", "POST content lưu nguyên khối trong transaction.", "Tùy chọn upload audio; PUT publish; refresh danh sách mine."]),
        ("7.3 Luyện tập/kiểm tra", ["GET published lessons → detail + assignments.", "Chọn toàn bài/section/exercise/assignment.", "Countdown và đồng bộ timeline; thu input piano.", "Client phân loại attempt; POST practice.", "Backend lưu session/attempt, cập nhật thống kê và hoàn thành assignment."]),
        ("7.4 Theo dõi và giao bổ sung", ["Teacher dashboard → danh sách học sinh.", "Mở tiến độ theo lesson và analytics.", "Xem biểu đồ, session, lỗi pitch/timing.", "POST assignment theo bài/section/exercise.", "Student thấy assignment trong lesson detail và luyện bằng assignmentId."]),
    ]
    for h, steps in workflows:
        add_heading(doc, h, 2)
        for s in steps: add_number(doc, s)

    add_heading(doc, "8. Ma trận truy vết", 1)
    trace = [
        ("SCR-01–05", "AUTH-01–04", "Auth APIs", "Users, PasswordResetTokens"),
        ("SCR-06", "LES-01, ASSESS-02", "Teacher dashboard, lessons/mine, students-progress", "Lessons, Users, PracticeSessions"),
        ("SCR-07–08", "MIDI-01, SCORE-01/02, EDIT-01/02, PED-01–03", "Lesson CRUD/content/publish", "Lessons, Notes, Sections, Annotations, Exercises"),
        ("SCR-09–10", "ASSESS-02, LEARN-03", "Teacher progress/analytics/assignments", "Users, Sessions, Attempts, Assignments"),
        ("SCR-11–12", "LES-01, SCORE-02, PED-01–03, LEARN-03", "Lessons + assignments", "Lessons và authoring tables, Assignments"),
        ("SCR-13–14", "LEARN-01/02, ASSESS-01", "Lesson detail + POST practice", "Notes, Sessions, Attempts, Progresses"),
        ("SCR-15", "ASSESS-02", "Practice history", "Sessions, Attempts, Lessons"),
    ]
    add_table(doc, ["Màn hình", "Feature", "API", "Table"], trace, [1500, 2800, 2600, 2460], font_size=8.4)

    add_heading(doc, "9. Yêu cầu phi chức năng", 1)
    nfr = [
        ("NFR-01", "Hiệu năng", "Danh sách lesson không tạo Cartesian explosion; detail tải theo nhu cầu; UI không block khi parse MIDI."),
        ("NFR-02", "Khả dụng", "Loading/error/retry/empty đầy đủ; API timeout có thông báo Việt hóa."),
        ("NFR-03", "Responsive", "Không cuộn ngang bản nhạc mobile; touch pan-y; layout thích ứng breakpoint."),
        ("NFR-04", "Nhất quán", "Teacher/student dùng cùng source và MusicXML pipeline."),
        ("NFR-05", "Bảo mật", "JWT, BCrypt, HMAC OTP, rate limit, secret ngoài source; ownership server-side."),
        ("NFR-06", "Dữ liệu", "UTC trong DB, UTC+7 tại UI; transaction khi thay content; migration có rollback."),
        ("NFR-07", "Khả năng bảo trì", "DTO tách entity, ProblemDetails, logging có correlation ID, API versioning."),
        ("NFR-08", "Đa nền tảng", "Web, Windows, Android/iOS; renderer adapter theo platform."),
    ]
    add_table(doc, ["ID", "Nhóm", "Tiêu chí"], nfr, [1100, 1600, 6660], font_size=9)

    add_heading(doc, "10. Bảo mật và quyền riêng tư", 1)
    add_heading(doc, "10.1 Kiểm soát hiện có", 2)
    for text in ["BCrypt cho mật khẩu; JWT Bearer chứa role và identity.", "OTP lưu HMAC-SHA256, hết hạn 10 phút, dùng một lần.", "Forgot password giới hạn 5 request/IP/phút và không tiết lộ email tồn tại.", "Token phía client lưu bằng flutter_secure_storage."]:
        add_bullet(doc, text)
    add_heading(doc, "10.2 Khoảng trống ưu tiên", 2)
    for text in ["Bắt buộc ownership check cho update/delete/publish/content/note/audio và analytics.", "GET lesson detail chỉ public với bài đã xuất bản; bản nháp chỉ owner.", "Không cho tự đăng ký role Teacher nếu chưa có invitation/admin approval.", "Backend phải tự tính judge result từ note/time, không tin hoàn toàn dữ liệu client.", "Whitelist MIME/extension, giới hạn dung lượng và dọn file audio khi xóa bài.", "Whitelist CORS theo môi trường; rate limit login/register/reset; audit log hành vi nhạy cảm.", "Đưa Gmail app password, JWT key và DB connection vào environment/user secrets; xoay vòng credential đã từng chia sẻ."]:
        add_bullet(doc, text)
    add_callout(doc, "Không lưu secret trong tài liệu", "Tài liệu này cố ý không chứa app password Gmail, JWT signing key hoặc connection string. Tài liệu vận hành chỉ được ghi tên biến môi trường và quy trình xoay vòng.", fill="FDECEC", accent=RED)

    add_heading(doc, "11. Chiến lược kiểm thử", 1)
    add_table(doc, ["Tầng", "Phạm vi", "Tiêu chí trọng yếu"], [
        ("Unit", "MIDI parser, MusicXML, judge, VietnamTime, validation", "Metadata gốc, chords/ties/rests, boundary timing, UTC+7"),
        ("API integration", "Auth, lesson content/delete, practice, analytics", "Status code, ownership, transaction, cascade/restrict"),
        ("Database", "Migration/rollback/index/delete", "Không schema drift; không orphan; query plan ổn định"),
        ("Widget", "Form, provider states, responsive", "Loading/error/empty/retry, focus guard piano"),
        ("Golden/visual", "Score teacher/student", "Cùng MusicXML cho cùng dữ liệu; không overlap"),
        ("Cross-platform", "Web/Windows/mobile", "OSMD/WebView, scroll dọc, touch, playback"),
        ("E2E", "Teacher import→publish→student practice→analytics", "Dữ liệu đi hết chuỗi và assignment hoàn thành"),
        ("Security", "Authz, IDOR, rate limit, upload", "Không truy cập chéo teacher; không lộ draft/stack trace"),
    ], [1500, 3150, 4710], font_size=8.7)
    add_heading(doc, "11.1 Tiêu chí nghiệm thu release", 2)
    for text in ["Import MIDI giữ đúng time signature, key signature, tempo map và track.", "Teacher/student render giống nhau; hợp âm có shared stem và nốt không chồng ngoài quy tắc ký âm.", "Mobile xem hết bản nhạc bằng cuộn dọc, không cần kéo ngang.", "Xóa lesson xử lý assignment/session/attempt/authoring data không lỗi FK.", "Practice section/exercise/assignment thực sự khởi chạy và lưu kết quả.", "Thời gian luyện/kiểm tra hiển thị đúng UTC+7.", "Bài bổ sung xuất hiện ở student và hoàn thành sau phiên có assignmentId.", "Không endpoint teacher nào thao tác tài nguyên không sở hữu."]:
        add_bullet(doc, text)

    add_heading(doc, "12. Quản lý triển khai dự án", 1)
    add_heading(doc, "12.1 Work Breakdown Structure", 2)
    add_table(doc, ["Workstream", "Deliverable", "Owner đề xuất", "Gate hoàn thành"], [
        ("WS1 Product/BA", "Scope, stories, rules, acceptance matrix", "PO/BA", "Sign-off phạm vi"),
        ("WS2 UX", "Responsive flow, editor interaction, state design", "UI/UX", "Prototype review"),
        ("WS3 Flutter", "Screens, MIDI/MusicXML, OSMD, realtime input", "Flutter Lead", "Analyze/test/cross-platform"),
        ("WS4 Backend", "Auth, lesson, practice, analytics, ownership", "Backend Lead", "API integration/security pass"),
        ("WS5 Database", "Schema, migration, indexes, retention", "Backend/DBA", "Migration + rollback rehearsal"),
        ("WS6 QA", "Test plan, automation, regression, UAT", "QA Lead", "Release acceptance passed"),
        ("WS7 DevOps", "Secrets, deploy, logging, monitoring, backup", "DevOps", "Production readiness review"),
    ], [1800, 3400, 1800, 2360], font_size=8.5)
    add_heading(doc, "12.2 Lộ trình ưu tiên", 2)
    add_table(doc, ["Giai đoạn", "Mục tiêu", "Hạng mục"], [
        ("P0 - Ổn định", "An toàn dữ liệu và quyền", "Ownership, draft visibility, schema drift, ProblemDetails, secret rotation"),
        ("P1 - Hiệu năng", "Danh sách và score tải nhanh", "Summary DTO, pagination, detail lazy-load, remove dead files"),
        ("P2 - Hoàn thiện nghiệp vụ", "Khép kín các luồng dở", "Audio UI, student progress, due date, teacher-student/class model"),
        ("P3 - Tin cậy đánh giá", "Chấm điểm server-authoritative", "Server judge, anti-duplicate, section/exercise session linkage"),
        ("P4 - Scale/Operations", "Vận hành production", "Health check, logging, metrics, audit, backup, offline renderer strategy"),
    ], [1600, 2600, 5160], font_size=8.8)
    add_heading(doc, "12.3 Definition of Done", 2)
    for text in ["Code review và không ghi đè thay đổi ngoài phạm vi.", "Migration có script/rollback và chạy trên bản sao dữ liệu.", "Unit/integration/widget/E2E liên quan đều đạt.", "Flutter analyze và backend build không có error mới.", "Đã kiểm tra Web, Windows và ít nhất một mobile target.", "API contract/tài liệu này được cập nhật; không có secret trong source/doc.", "Có tiêu chí observability và kế hoạch rollback release."]:
        add_bullet(doc, text)

    add_heading(doc, "13. Risk register và technical debt", 1)
    risks = [
        ("R01", "Critical", "IDOR/thiếu ownership ở lesson và analytics", "Truy cập/sửa dữ liệu chéo", "Filter teacher ownership + class relation + tests"),
        ("R02", "High", "GET lesson detail public lộ draft", "Rò rỉ nội dung", "Published-or-owner policy"),
        ("R03", "High", "Client tự gửi judge result", "Gian lận/sai thống kê", "Server-side judge"),
        ("R04", "High", "DbContext/snapshot lệch delete behavior", "Migration/xóa bài lỗi", "Migration đồng bộ + DB tests"),
        ("R05", "High", "Lesson list trả toàn bộ notes", "Payload/timeout lớn", "Summary DTO + pagination"),
        ("R06", "High", "Save content thay note và xóa attempt cũ", "Mất lịch sử đánh giá", "Lesson revision/versioning"),
        ("R07", "Medium", "OSMD phụ thuộc CDN", "Offline/blocked network không render", "Bundle asset hoặc fallback"),
        ("R08", "Medium", "Audio upload không validate/dọn file", "Bảo mật và file mồ côi", "Limit/MIME/cleanup"),
        ("R09", "Medium", "Router không chặn route theo role", "UX và defense-in-depth yếu", "Role route guard"),
        ("R10", "Medium", "nvarchar(max) dùng rộng", "Dung lượng/query/index kém", "MaxLength + migration"),
        ("R11", "Medium", "Dashboard scope lẫn global/teacher", "KPI sai nghĩa", "Thống nhất scope"),
        ("R12", "Low", "Placeholder/dead files", "Nợ bảo trì", "Xóa hoặc lập backlog hoàn thiện"),
    ]
    add_table(doc, ["ID", "Mức", "Rủi ro", "Tác động", "Ứng phó"], risks, [650, 900, 2850, 2150, 2810], font_size=8.1)

    add_heading(doc, "14. Triển khai và cấu hình", 1)
    add_table(doc, ["Hạng mục", "Hiện tại", "Khuyến nghị production"], [
        ("API URL", "localhost:5238 / 10.0.2.2 emulator", "Build flavor/env; HTTPS domain"),
        ("Database", "SQL Server, auto Migrate + Seed startup", "Migration job riêng; backup trước deploy"),
        ("Secrets", "Configuration", "Environment/user-secrets/vault; rotation"),
        ("CORS", "AllowAnyOrigin/Header/Method", "Whitelist origin"),
        ("Audio", "wwwroot/uploads/audio", "Object storage + signed URL + lifecycle"),
        ("OSMD", "jsDelivr CDN 1.9.9", "Pin checksum/bundle local/fallback"),
        ("Monitoring", "Chưa chuẩn hóa", "Structured logs, metrics, traces, health endpoint"),
    ], [1800, 3100, 4460], font_size=8.8)
    add_heading(doc, "14.1 Biến cấu hình cần quản lý", 2)
    for text in ["ConnectionStrings__DB", "Jwt__Key / Jwt__Issuer / Jwt__Audience", "Gmail__Username / Gmail__AppPassword / Gmail__FromName", "AllowedOrigins", "Uploads__MaxSize / AllowedMimeTypes", "OSMD asset/CDN version"]:
        add_bullet(doc, text)

    add_heading(doc, "15. Phụ lục", 1)
    add_heading(doc, "15.1 Lịch sử migration", 2)
    migrations = ["InitialMusicIBookDb", "UpdatePracticeAndNoteModels", "AddGoogleLoginSupport", "AddJudgeResultToAttempt", "PreserveMidiNotationTiming", "AddPolyphonicScoreMetadata", "AddTeacherAuthoringWorkflow", "OptimizeDatabase", "PreserveMidiTracks"]
    for i, m in enumerate(migrations, 1): add_number(doc, f"{m}")
    add_heading(doc, "15.2 Thuật ngữ", 2)
    add_table(doc, ["Thuật ngữ", "Định nghĩa"], [
        ("Beat", "Vị trí thời gian âm nhạc theo phách, độc lập giây thực."), ("Staff", "Khuông nhạc; thường map tay phải/tay trái."),
        ("Voice", "Dòng giai điệu độc lập trong cùng khuông."), ("MusicXML", "Định dạng trao đổi ký âm dùng để render bản nhạc."),
        ("OSMD", "OpenSheetMusicDisplay, renderer MusicXML thành SVG."), ("Assignment", "Bài bổ sung giao riêng cho một học sinh."),
        ("Attempt", "Một kết quả chơi của học sinh so với nốt chuẩn."), ("Schema drift", "Khác biệt giữa model runtime, snapshot/migration và database thực."),
    ], [2200, 7160], font_size=9)
    add_heading(doc, "15.3 Cơ sở và giới hạn tài liệu", 2)
    add_body(doc, "Tài liệu được lập từ mã nguồn trong workspace tại thời điểm baseline. Nội dung 'đã triển khai' phản ánh route/service/model hiện hữu, không thay thế UAT trên môi trường production. Các đề xuất bảo mật, scale và vận hành được tách rõ khỏi hành vi as-is.")

    # Core properties and save.
    doc.core_properties.title = "Tài liệu quản lý dự án Music IBook"
    doc.core_properties.subject = "Database, screens, features, APIs, workflows and project governance"
    doc.core_properties.author = "Music IBook Project Team"
    doc.core_properties.keywords = "Music IBook, Flutter, ASP.NET Core, SQL Server, MusicXML, OSMD"
    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build_document())
